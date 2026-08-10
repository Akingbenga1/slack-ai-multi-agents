"""Unit tests for document parsers (Task 8.1)."""

from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from api.app.ingest.documents import (
    DocumentFormat,
    DocumentUnitKind,
    UnsupportedDocumentFormatError,
    detect_document_format,
    extract_csv,
    extract_document,
    extract_docx,
    extract_pdf,
    extract_xlsx,
)


def _pdf_with_text(text: str) -> bytes:
    """Build a minimal one-page PDF that embeds extractable text."""
    # Use reportlab-free approach: pypdf cannot easily add text; write a tiny
    # PDF with a content stream containing the string (Type1 fonts).
    # Simpler: write via a content stream using Helvetica Tj operator.
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    # Build a content stream: BT /F1 12 Tf 72 720 Td (text) Tj ET
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    content = f"BT /F1 12 Tf 72 720 Td ({safe}) Tj ET\n".encode("latin-1", errors="replace")

    stream = DecodedStreamObject()
    stream.set_data(content)
    stream_ref = writer._add_object(stream)

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)

    resources = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref}),
        }
    )
    page[NameObject("/Resources")] = resources
    page[NameObject("/Contents")] = stream_ref

    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _docx_bytes(*, paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    doc = Document()
    doc.core_properties.title = "Fixture Doc"
    for p in paragraphs:
        doc.add_paragraph(p)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r_i, row in enumerate(table_rows):
            for c_i, cell in enumerate(row):
                table.rows[r_i].cells[c_i].text = cell
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[object]], *, sheet_name: str = "Sheet1") -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_detect_format_from_extension_and_mime():
    assert detect_document_format(filename="a.PDF") is DocumentFormat.PDF
    assert detect_document_format(content_type="text/csv; charset=utf-8") is DocumentFormat.CSV
    assert detect_document_format(format="docx") is DocumentFormat.DOCX
    with pytest.raises(UnsupportedDocumentFormatError):
        detect_document_format(filename="notes.txt")


def test_extract_pdf_pages():
    raw = _pdf_with_text("Onboarding checklist for new hires")
    result = extract_pdf(raw, filename="guide.pdf")
    assert result.source_format is DocumentFormat.PDF
    assert result.filename == "guide.pdf"
    assert len(result.units) >= 1
    assert result.units[0].kind is DocumentUnitKind.PAGE
    assert result.units[0].locator == "page=1"
    assert "Onboarding checklist" in result.text


def test_extract_docx_paragraphs_and_table():
    raw = _docx_bytes(
        paragraphs=["Welcome to Acme", ""],
        table_rows=[["Step", "Owner"], ["Invite Slack", "IT"]],
    )
    result = extract_docx(raw, filename="playbook.docx")
    assert result.title == "Fixture Doc"
    kinds = [u.kind for u in result.units]
    assert DocumentUnitKind.PARAGRAPH in kinds
    assert DocumentUnitKind.TABLE in kinds
    assert any("Welcome to Acme" in u.text for u in result.units)
    assert any("Invite Slack" in u.text for u in result.units)


def test_extract_csv_rows():
    csv = "product,price\nWidget,9.99\nGadget,12.50\n"
    result = extract_csv(csv, filename="catalog.csv")
    assert len(result.units) == 2
    assert result.units[0].locator == "row=1"
    assert "product: Widget" in result.units[0].text
    assert "price: 9.99" in result.units[0].text


def test_extract_xlsx_all_sheets():
    raw = _xlsx_bytes(
        [["sku", "qty"], ["A-1", 3]],
        sheet_name="Inventory",
    )
    result = extract_xlsx(raw, filename="stock.xlsx")
    assert len(result.units) == 1
    assert result.units[0].kind is DocumentUnitKind.TABLE_ROW
    assert "sheet=Inventory!row=1" == result.units[0].locator
    assert "sku: A-1" in result.units[0].text


def test_extract_document_dispatch(tmp_path: Path):
    path = tmp_path / "notes.csv"
    path.write_text("title,body\nHello,World\n", encoding="utf-8")
    result = extract_document(path)
    assert result.source_format is DocumentFormat.CSV
    assert result.filename == "notes.csv"
    assert "title: Hello" in result.text


def test_empty_csv_yields_no_units():
    result = extract_csv("a,b\n,\n", filename="empty.csv")
    assert result.units == []
